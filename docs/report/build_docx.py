#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把合并后的 Markdown 报告排版进学校 Word 模板。

用法：
    python docs/report/build_docx.py
    python docs/report/build_docx.py --md docs/report/DecisionJury_小学期实验报告.md \
        --out docs/report/DecisionJury_小学期实验报告.docx

说明：
- 脚本只读取 `docs/report/templates/小学期实验报告模板.docx`（学校模板的只读副本），
  绝不写回模板文件；原始模板 `小学期实验报告模板(1).docx` 不会被修改。
- 模板中“说明：指导教师评分后……”之后的内容会被替换为报告正文。
- 图片路径相对于 docs/report/ 解析。
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = BASE_DIR / "templates" / "小学期实验报告模板.docx"
DEFAULT_MD = BASE_DIR / "DecisionJury_小学期实验报告.md"
DEFAULT_OUT = BASE_DIR / "DecisionJury_小学期实验报告.docx"

BODY_EAST = "宋体"
BODY_ASCII = "Times New Roman"
CODE_ASCII = "Consolas"
BODY_SIZE = 10.5
CODE_SIZE = 9


def set_run_font(run, size=BODY_SIZE, bold=False, ascii_font=BODY_ASCII, east_font=BODY_EAST, color=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_font)


def shade_paragraph(paragraph, fill="F2F2F2"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill="D9E2F3"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_paragraph_text(paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[(.*?)\]\((.*?)\)", r"\1", text)
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = text.replace("`", "")
    text = text.replace("<br>", " ").replace("<br/>", " ")
    return text.strip()


def add_paragraph(doc: Document, text: str, size=BODY_SIZE, bold=False, indent=True,
                  align=None, space_after=6, line_spacing=1.5, color=None):
    p = doc.add_paragraph()
    run = p.add_run(clean_inline(text))
    set_run_font(run, size=size, bold=bold, color=color)
    pf = p.paragraph_format
    if indent:
        pf.first_line_indent = Pt(21)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc: Document, text: str, level: int):
    # 学校模板要求标题为五号宋体加粗；这里统一用 10.5pt，避免改变模板字号约定。
    p = doc.add_paragraph()
    run = p.add_run(clean_inline(text))
    set_run_font(run, size=BODY_SIZE, bold=True)
    pf = p.paragraph_format
    pf.space_before = Pt(10 if level == 1 else 6)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    return p


def add_code_block(doc: Document, code: str):
    p = doc.add_paragraph()
    run = p.add_run(code)
    set_run_font(run, size=CODE_SIZE, ascii_font=CODE_ASCII, east_font=BODY_EAST)
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(2)
    pf.line_spacing = 1.05
    shade_paragraph(p, "F7F7F7")
    return p


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(clean_inline(row[ci]) if ci < len(row) else "")
            set_run_font(run, size=9, bold=(ri == 0))
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if ri == 0:
                shade_cell(cell, "D9E2F3")
    return table


def add_image(doc: Document, alt: str, rel_path: str):
    path = (BASE_DIR / rel_path).resolve()
    if not path.exists():
        add_paragraph(doc, f"【图片缺失：{rel_path}】", indent=False, color=RGBColor(0xC0, 0x00, 0x00))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(alt)
    set_run_font(cap_run, size=9, bold=False)
    cap.paragraph_format.space_after = Pt(8)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def iter_blocks(lines: list[str]):
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip("\n")
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("<!--"):
            while i < n and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            yield ("code", "\n".join(code_lines))
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            yield ("table", parse_table(table_lines))
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if m:
            yield ("image", (m.group(1), m.group(2)))
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            yield ("heading", (level, stripped.lstrip("#").strip()))
            i += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            yield ("quote", "\n".join(quote_lines))
            continue
        if re.match(r"^\s*[-*]\s+", line):
            items = []
            while i < n and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip())
                i += 1
            yield ("bullets", items)
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < n and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]).strip())
                i += 1
            yield ("numbers", items)
            continue
        if stripped == "---":
            i += 1
            continue
        para_lines = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt.startswith(("#", "|", "```", ">", "![", "---", "<!--")):
                break
            if re.match(r"^\s*[-*]\s+", lines[i]) or re.match(r"^\s*\d+\.\s+", lines[i]):
                break
            para_lines.append(nxt)
            i += 1
        yield ("paragraph", " ".join(para_lines))


def fill_cover(doc: Document) -> None:
    """填写封面字段；只处理成绩页之前的段落。"""
    cover_end = None
    for idx, p in enumerate(doc.paragraphs):
        if "教务部" in p.text or "实践报告成绩" in p.text:
            cover_end = idx
            break
    if cover_end is None:
        cover_end = len(doc.paragraphs)

    replacements = [
        ("课程名称", "课程名称 ：      综合能力实训"),
        ("项目名称", "项目名称 ：      DecisionJury——基于多 Agent 协作的日常冷静决策助手"),
        ("班级", "班    级 ：      【请填写：班级】"),
        ("专业", "专    业 ：      【请填写：专业】"),
        ("任课教师", "任课教师 ：     魏振文、赵双华"),
        ("学号", "学     号：      【请填写：全组学号】"),
        ("姓名", "姓     名：      【请填写：全组成员姓名】"),
    ]
    for idx in range(cover_end):
        text = doc.paragraphs[idx].text
        compact = text.replace(" ", "")
        for key, value in replacements:
            if key in compact:
                set_paragraph_text(doc.paragraphs[idx], value)
                break


def remove_template_sample(doc: Document) -> None:
    """保留封面和成绩页，删除“说明：指导教师评分后……”之后的模板示例正文。"""
    body = doc.element.body
    note_par = None
    for p in doc.paragraphs:
        if "指导教师评分后" in p.text:
            note_par = p
            break
    if note_par is None:
        return
    removing = False
    for child in list(body):
        if child is note_par._element:
            removing = True
            continue
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def build_docx(md_path: Path, out_path: Path, template_path: Path) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f"缺少模板文件：{template_path}")
    if not md_path.exists():
        raise FileNotFoundError(f"缺少合并报告：{md_path}")
    # 安全保护：只读模板，绝不写回模板文件。
    if out_path.resolve() == template_path.resolve():
        raise SystemExit("输出文件不能与模板文件相同，请通过 --out 指定其他路径。")
    if template_path.resolve().parent in out_path.resolve().parents:
        raise SystemExit("输出文件不能写入 templates/ 目录，以免覆盖模板副本。")

    text = md_path.read_text(encoding="utf-8")
    # 从 "# 00" 章节开始，跳过合并文件头部说明
    start = text.find("# 00 ")
    if start == -1:
        start = 0
    lines = text[start:].splitlines()

    doc = Document(str(template_path))
    fill_cover(doc)
    remove_template_sample(doc)

    doc.add_page_break()

    skip_until_heading = False
    first_h1 = True
    for kind, payload in iter_blocks(lines):
        if kind == "heading":
            level, title = payload
            if "封面信息" in title:
                skip_until_heading = True
                continue
            if skip_until_heading:
                if level == 2:
                    skip_until_heading = False
                else:
                    continue
            if level == 1:
                if not first_h1:
                    doc.add_page_break()
                first_h1 = False
            add_heading(doc, title, level)
        elif skip_until_heading:
            continue
        elif kind == "quote":
            first_line = payload.splitlines()[0] if payload else ""
            if first_line.startswith(("负责人：", "状态：", "说明：")):
                continue
            add_paragraph(doc, payload.replace("\n", " "), size=9, indent=False,
                          color=RGBColor(0x55, 0x55, 0x55))
        elif kind == "paragraph":
            add_paragraph(doc, payload)
        elif kind == "bullets":
            for item in payload:
                p = add_paragraph(doc, "• " + item, indent=False, space_after=3)
                p.paragraph_format.left_indent = Pt(21)
        elif kind == "numbers":
            for idx, item in enumerate(payload, start=1):
                p = add_paragraph(doc, f"{idx}. " + item, indent=False, space_after=3)
                p.paragraph_format.left_indent = Pt(21)
        elif kind == "code":
            add_code_block(doc, payload)
        elif kind == "table":
            add_table(doc, payload)
            add_paragraph(doc, "", indent=False, space_after=2)
        elif kind == "image":
            alt, rel_path = payload
            add_image(doc, alt, rel_path)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"已生成 Word 报告：{out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="生成小学期实验报告 Word 文档")
    parser.add_argument("--md", default=str(DEFAULT_MD), help="合并后的 Markdown 报告路径")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="输出 docx 路径")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Word 模板路径")
    args = parser.parse_args()
    build_docx(Path(args.md).resolve(), Path(args.out).resolve(), Path(args.template).resolve())


if __name__ == "__main__":
    main()
